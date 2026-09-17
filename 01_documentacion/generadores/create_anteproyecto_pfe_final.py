import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

def generate_doc():
    doc = Document()

    # Configuración de márgenes para 2 páginas exactas
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    NAVY = RGBColor(27, 54, 93)       # #1B365D - Títulos
    STEEL = RGBColor(70, 130, 180)    # #4682B4 - Subtítulos/Énfasis
    CHARCOAL = RGBColor(40, 40, 40)   # #282828 - Texto general

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.font.color.rgb = CHARCOAL

    # Encabezado
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run("Anteproyecto de Proyecto Final de Estudios (PFE)")
    r_title.font.size = Pt(13.5)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(8)
    r_sub = sub_p.add_run("Carrera de Ingeniería Mecatrónica")
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    def add_section(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return p

    # 1. Información General
    add_section("1. Información General")
    info_items = [
        ("Título:", " Sistema de Visión Artificial y Control Co-simulado Adaptativo para la Pulverización Selectiva de Oídio (Erysiphe necator / Oidium tuckeri) en Vid"),
        ("Alumno:", " Nicolás Ezequiel Sorrentino"),
        ("Foco Declarado:", " Sistemas Simulados")
    ]
    for tag, val in info_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.12
        r_tag = p.add_run(tag)
        r_tag.bold = True
        r_tag.font.color.rgb = NAVY
        p.add_run(val)

    # 2. Fundamentación y Problema
    add_section("2. Fundamentación y Problema")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "En la producción vitivinícola actual, el tratamiento contra el oídio de la vid se realiza predominantemente de forma "
        "preventiva y continua sobre la totalidad del espaldero. Esta práctica tradicional resulta ineficiente cuando los focos "
        "de infección son tempranos o aislados, ya que provoca un desperdicio significativo de producto químico, incrementa los costos "
        "operativos y genera un impacto ambiental negativo sobre el suelo y el entorno del viñedo."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Desarrollar una solución de pulverización selectiva eficiente implica resolver dos desafíos clave desde la ingeniería mecatrónica:"
    )

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r_bold = p.add_run("Percepción visual a campo: ")
    r_bold.bold = True
    p.add_run(
        "Lograr una detección automatizada y confiable en condiciones reales de viñedo (iluminación solar variable, sombras, "
        "hojas solapadas y fondos con vegetación densa). El sistema debe identificar con precisión las hojas que presentan infección "
        "por oídio y diferenciarlas claramente de las hojas sanas, operando en tiempo real para acompañar el avance del vehículo."
    )

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    r_bold = p.add_run("Estrategia de pulverización inteligente: ")
    r_bold.bold = True
    p.add_run(
        "Aplicar el fitosanitario únicamente sobre la hoja visible detectada es insuficiente, debido a que el hongo disemina sus "
        "esporas por acción del viento hacia la vegetación vecina antes de que los síntomas sean visibles. Por ello, el sistema "
        "debe calcular una zona de seguridad alrededor de las hojas enfermas detectadas, adaptando el ancho del rociado según el "
        "viento y las condiciones climáticas del momento."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "El problema que aborda este proyecto consiste en integrar ambas etapas en un lazo cerrado: entrenar un modelo de visión por "
        "computadora para la detección de hojas infectadas y conectarlo con un sistema de control que sincronice el avance del vehículo "
        "y accione selectivamente una barra de boquillas para tratar los focos de manera oportuna y económica."
    )

    # 3. Objetivos del Proyecto
    add_section("3. Objetivos del Proyecto")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_obj_g = p.add_run("Objetivo General:")
    r_obj_g.bold = True
    r_obj_g.font.color.rgb = NAVY
    p.add_run(
        " Desarrollar y validar mediante co-simulación un sistema mecatrónico que integre visión artificial y control adaptativo "
        "para la pulverización selectiva de oídio en vid, optimizando el uso de agroquímicos."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_obj_e = p.add_run("Objetivos Específicos:")
    r_obj_e.bold = True
    r_obj_e.font.color.rgb = NAVY

    objetivos_esp = [
        "Compilar y preparar un conjunto de imágenes de viñedos a campo real para entrenar un modelo de visión artificial capaz de detectar hojas con oídio.",
        "Diseñar una lógica de control que calcule una zona de pulverización ampliada alrededor de las hojas infectadas, adaptándose a variables meteorológicas (viento y humedad).",
        "Modelar la sincronización temporal entre el avance del vehículo y la activación de las boquillas, compensando la distancia física entre la cámara y la barra pulverizadora.",
        "Implementar un entorno de co-simulación para evaluar el desempeño global del sistema en lazo cerrado y cuantificar el porcentaje de ahorro de producto frente a la aplicación continua."
    ]
    for obj in objetivos_esp:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.12
        p.add_run(obj)

    # 4. Metodología Propuesta
    add_section("4. Metodología Propuesta")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "El desarrollo del proyecto se llevará a cabo a través de tres etapas metodológicas integradas:"
    )

    etapas = [
        ("1. Modelo de Percepción Visual: ",
         "Se recopilarán y curarán imágenes representativas de viñedos adultos a campo abierto. Se entrenará un modelo de aprendizaje "
         "profundo para detección de objetos (familia YOLO), enfocado en localizar mediante cajas delimitadoras las hojas que presenten "
         "síntomas de oídio. Se priorizará un equilibrio entre precisión y velocidad de procesamiento para permitir su operación en tiempo real."),

        ("2. Zona de Pulverización Adaptativa: ",
         "Se formulará un algoritmo que tome la ubicación de las hojas enfermas detectadas y consulte datos meteorológicos locales (velocidad "
         "del viento y humedad). Con esta información se determinará una zona de cobertura de seguridad alrededor del foco para neutralizar "
         "las esporas que pudieran haberse diseminado hacia la vegetación circundante."),

        ("3. Co-simulación y Evaluación de Desempeño: ",
         "Se desarrollará un entorno de simulación que reproduzca el avance del vehículo a lo largo de la hilera y la apertura/cierre de "
         "las electroválvulas en una barra vertical. La simulación integrará la detección de la cámara, coordinará el momento exacto de "
         "disparo cuando la barra alcance la posición de las hojas infectadas y medirá la reducción en el consumo de fitosanitario.")
    ]
    for tag, val in etapas:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        r_tag = p.add_run(tag)
        r_tag.bold = True
        p.add_run(val)

    # 5. Diseño Preliminar del Sistema
    add_section("5. Diseño Preliminar del Sistema")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "La arquitectura del sistema propuesto se estructura en tres subsistemas interconectados:"
    )

    subs = [
        ("Subsistema de Percepción: ",
         "Adquiere las imágenes de la vegetación mediante una cámara frontal. El modelo de visión detecta las hojas con oídio y traduce "
         "su posición en la imagen a coordenadas espaciales respecto a la altura del espaldero y el avance del vehículo."),

        ("Subsistema de Planificación y Control: ",
         "Recibe las coordenadas de las hojas infectadas, define la zona de pulverización requerida según las condiciones meteorológicas y "
         "calcula el tiempo de espera necesario (según la velocidad del móvil) para activar las boquillas justo cuando pasen frente al área a tratar."),

        ("Subsistema Físico Simulado (Planta y Boquillas): ",
         "Modela el comportamiento de una barra vertical con boquillas independientes a diferentes alturas, simulando la respuesta de las "
         "electroválvulas y la aplicación del fluido sobre la vegetación para validar la cobertura del tratamiento.")
    ]
    for tag, val in subs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        r_tag = p.add_run(tag)
        r_tag.bold = True
        p.add_run(val)

       # Guardar en 01_documentacion/anteproyecto
    out_path = Path(__file__).resolve().parent.parent / "anteproyecto" / "Anteproyecto_PFE_Sorrentino.docx"
    doc.save(str(out_path))
    print(f"Documento generado exitosamente en: {out_path.resolve()}")

if __name__ == "__main__":
    generate_doc()
